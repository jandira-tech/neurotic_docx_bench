# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
# SPDX-License-Identifier: AGPL-3.0-only
"""Generate the fixed, complex starting document from reproducible source."""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw


def build_seed(destination: Path) -> None:
    doc = Document()
    doc.core_properties.title = "Jubarte Live Bench — fixed starting point"
    doc.core_properties.author = "Jandira Technologies, LLC"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 8)
    doc.styles["Normal"].font.name = "Liberation Serif"
    doc.styles["Normal"].font.size = Pt(11)
    section = doc.sections[0]
    section.header.paragraphs[0].text = "JUBARTE / LIVE BENCH / REFERENCE AGREEMENT"
    footer = section.footer.paragraphs[0]
    footer.add_run("Confidential • Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_heading("Research Collaboration Agreement", 0)
    p = doc.add_paragraph(
        "This agreement establishes the responsibilities of the parties, effective 8 October 2026."
    )
    doc.add_comment(
        p.runs, text="Review the effective date and responsibilities.", author="Jubarte Bench", initials="JB"
    )
    doc.add_heading("1. Scope and definitions", 1)
    for text in (
        "Deliverables include reports, annotated documents, supporting evidence and revisions.",
        "A material change affects the scope, schedule, cost, intellectual property or confidentiality.",
        "Acceptance requires written approval and a complete record of the review.",
    ):
        doc.add_paragraph(text, style="List Number")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    for cell, label in zip(
        table.rows[0].cells, ("Milestone", "Owner", "Due date", "Acceptance"), strict=True
    ):
        cell.text = label
    for row in (
        ("Discovery", "Research team", "15 October", "Approved scope"),
        ("Prototype", "Engineering", "1 November", "Reproducible evidence"),
        ("Final report", "Joint committee", "15 November", "Signed decision"),
    ):
        for cell, text in zip(table.add_row().cells, row, strict=True):
            cell.text = text
    nested = table.rows[2].cells[3].add_table(rows=2, cols=2)
    nested.cell(0, 0).text, nested.cell(0, 1).text = "Check", "Result"
    nested.cell(1, 0).text, nested.cell(1, 1).text = "Source integrity", "Required"
    doc.add_heading("2. Review procedure", 1)
    p = doc.add_paragraph()
    p.add_run("Proposed changes ").bold = True
    p.add_run("must remain traceable to their source. ").italic = True
    p.add_run("Rejected changes retain the original text.").font.color.rgb = RGBColor(140, 40, 40)
    image = Image.new("RGB", (640, 160), "white")
    draw = ImageDraw.Draw(image)
    for i, label in enumerate(("SOURCE", "COMPARE", "REVIEW", "PUBLISH")):
        x = i * 160
        draw.rectangle((x + 10, 40, x + 140, 120), outline=(30, 90, 120), width=3)
        draw.text((x + 30, 75), label, fill=(30, 90, 120))
    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    doc.add_picture(stream, width=Inches(6))
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading("Schedule A — detailed terms", 1)
    for i in range(1, 9):
        doc.add_heading(f"A.{i} Obligations and exceptions", 2)
        doc.add_paragraph(
            f"The parties shall maintain complete evidence for milestone {i}. Any exception must "
            "identify its source, responsible owner, date, reason and effect on the acceptance criteria."
        )
        doc.add_paragraph(
            "• Preserve the original document.\n• Record every failure.\n• Verify before publication."
        )
    doc.add_paragraph("Authorized signatures: ____________________     ____________________")
    buffer = BytesIO()
    doc.save(buffer)
    # Normalize ZIP metadata too: stable contents have a stable hash across rebuilds.
    destination.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(buffer) as source, ZipFile(destination, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, (2026, 10, 8, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            data = source.read(name)
            if name == "word/comments.xml":
                import re

                data = re.sub(rb'w:date="[^"]*"', b'w:date="2026-10-08T00:00:00Z"', data)
            target.writestr(entry, data)
