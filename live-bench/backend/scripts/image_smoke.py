# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
# SPDX-License-Identifier: AGPL-3.0-only
import json
import shutil
import sys
import uuid
from datetime import UTC, datetime
from pathlib import Path
from time import monotonic

import httpx
from docx import Document as Docx

from livebench.benchmark import Benchmark
from livebench.config import Config
from livebench.models import MAX_FILE_BYTES, Document, ToolFailure
from livebench.seed import build_seed
from livebench.tools import RealTools, command

root = Path("/tmp/check")
root.mkdir()
a, b = root / "a.docx", root / "b.docx"
build_seed(a)
d = Docx(a)
d.paragraphs[1].add_run(" Revised responsibilities and a new deadline.")
d.save(b)
work = root / "work"
work.mkdir()
shutil.copy(a, work / "previous.docx")
shutil.copy(b, work / "source.docx")
client = httpx.Client(
    transport=httpx.MockTransport(
        lambda _: httpx.Response(503, json={"error": "deliberate offline reference"})
    )
)
tools = RealTools(
    Config(root, "https://word.invalid", "dummy-word", "https://ingest.invalid", "dummy-ingest"), client
)
print(json.dumps({"versions": tools.versions()}), flush=True)
network_probe = (
    "import socket,os\nassert 'WORD_API_TOKEN' not in os.environ\n"
    "s=socket.socket(socket.AF_UNIX)\ns.close()\n"
    "try: socket.socket(socket.AF_INET)\n"
    "except PermissionError: print('Network isolated')\n"
    "else: raise RuntimeError('network permitted')"
)
command([sys.executable, "-c", network_probe], 10)
assert MAX_FILE_BYTES == 20_000_000
try:
    command(
        [
            sys.executable,
            "-c",
            "with open('/tmp/oversized-probe', 'wb') as fh:\n"
            " fh.write(b'x' * 20000000)\n fh.write(b'y')\n fh.flush()",
        ],
        10,
    )
except ToolFailure:
    pass
else:
    raise AssertionError("subprocess file size limit was not enforced")
result = Benchmark(tools, monotonic, lambda: datetime.now(UTC).isoformat()).run(
    str(uuid.uuid7()),
    1,
    100,
    Document("seed", "generated://starting_point.docx", a, "a"),
    Document("modified", "generated://test.docx", b, "b"),
    work,
    lambda p: print("phase=" + p, flush=True),
)
print(
    json.dumps(
        {"scores": result["scores"], "stages": result["stages"], "artifacts": len(result["artifacts"])}
    ),
    flush=True,
)
assert all(s["status"] in ("ok", "reference") for s in result["scores"]), result["scores"]
assert all(s["status"] == "ok" for s in result["stages"] if s["tool"] != "word"), result["stages"]
assert len(list(work.glob("score-*.png"))) == 4
assert all(artifact["size"] <= MAX_FILE_BYTES for artifact in result["artifacts"])
print("Production image tool/network/scoring smoke passed", flush=True)
